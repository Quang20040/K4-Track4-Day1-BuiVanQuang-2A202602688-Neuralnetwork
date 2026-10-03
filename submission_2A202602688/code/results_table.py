"""results_table.py — Quản lý lưu trữ kết quả và xuất bảng Excel experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result['cfg'], result['history'], result['summary'] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    r_dir = Path(results_dir)
    r_dir.mkdir(parents=True, exist_ok=True)

    exp_id = result["cfg"]["exp_id"]
    out_file = r_dir / f"{exp_id}.json"

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }

    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(out_file)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    r_dir = Path(results_dir)
    if not r_dir.exists():
        return []

    results = []
    for p in sorted(r_dir.glob("*.json")):
        with open(p, "r", encoding="utf-8") as f:
            results.append(json.load(f))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png".
    """
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg["exp_id"]

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", "sgd_momentum"),
        "lr": cfg.get("lr", 0.0),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": str(cfg.get("hidden", (256, 128))),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", ""),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_scores.get("accuracy") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("macro_f1") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes or cfg.get("notes", ""),
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước:
      1. wb = openpyxl.load_workbook(template_path)  # Giữ nguyên công thức
      2. ws = wb['Experiments']; đọc tiêu đề dòng 1 để map cột
      3. Ghi giá trị vào đúng cột; giữ nguyên hoặc sao chép công thức ở các cột công thức
      4. wb.save(out_path)
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    # Đọc headers ở hàng 1
    col_names = {}
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col_idx).value
        if val is not None:
            col_names[str(val).strip()] = col_idx

    formula_cols = {}
    # Phát hiện các cột có công thức từ dòng 2 (nếu có)
    for col_name, col_idx in col_names.items():
        cell_val = str(ws.cell(row=2, column=col_idx).value or "")
        if cell_val.startswith("="):
            formula_cols[col_idx] = cell_val

    # Điền từng dòng
    start_row = 2
    for r_i, r_data in enumerate(rows):
        cur_row = start_row + r_i
        for key, val in r_data.items():
            if key in col_names:
                col_idx = col_names[key]
                ws.cell(row=cur_row, column=col_idx, value=val)

        # Xử lý các cột công thức (cập nhật số hàng trong công thức)
        for col_idx, base_formula in formula_cols.items():
            # Thay đổi chỉ số hàng trong công thức từ 2 -> cur_row
            import re
            new_formula = re.sub(r"([A-Za-z]+)2", rf"\g<1>{cur_row}", base_formula)
            ws.cell(row=cur_row, column=col_idx, value=new_formula)

    out_file = Path(out_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_file)
    print(f"Đã cập nhật bảng thực nghiệm vào {out_file} ({len(rows)} thí nghiệm).")
